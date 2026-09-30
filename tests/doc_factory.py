"""Build real test documents in memory (no files on disk, no extra dependencies)."""

import io


def _pdf_escape(line: str) -> str:
    return line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(lines: list[str]) -> bytes:
    """A valid one-page PDF with real (extractable) Helvetica text."""
    ops = ["BT /F1 11 Tf 50 780 Td 13 TL"]
    ops += [f"({_pdf_escape(line)}) Tj T*" for line in lines]
    ops.append("ET")
    stream = "\n".join(ops).encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref,
    )
    return bytes(out)


def make_blank_pdf() -> bytes:
    """Like a scanned page: a valid PDF with no text layer."""
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def make_encrypted_pdf() -> bytes:
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(make_pdf(["secret notes"]))))
    writer.encrypt("pw", algorithm="RC4-128")
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    import docx

    document = docx.Document()
    for p in paragraphs:
        document.add_paragraph(p)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


PNG_BYTES = b"\x89PNG\r\n\x1a\n" + b"\x00" * 64
EXE_BYTES = b"MZ\x90\x00\x03" + b"\x00" * 64
